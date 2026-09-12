"""The index sheet: one row per thing that arrived, in CSV and XLSX.

CSV because it diffs, greps and imports. XLSX because it is what the client
actually double-clicks, and a filing job that ends in "now open this in
LibreOffice and it looks like a spreadsheet" lands differently from one that
ends in a CSV.

Both are byte-for-byte reproducible. A second run over an unchanged mailbox has
to leave the index unchanged, and the only way to *show* that rather than claim
it is for `cmp` to agree — which for XLSX means flattening the two timestamps a
zip archive and an Office document would otherwise stamp into every file.
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from datetime import datetime
from pathlib import Path

from .models import Action, Plan

COLUMNS = (
    "date",
    "sender",
    "subject",
    "original_filename",
    "new_path",
    "size_bytes",
    "rule",
    "status",
    "note",
    "source_message",
)

#: Frozen so that two runs produce identical bytes. openpyxl would otherwise
#: stamp `docProps/core.xml` with the current time, and a sheet that differs
#: every run cannot be used as evidence that nothing changed.
EPOCH = datetime(1980, 1, 1, 0, 0, 0)


def rows(plan: Plan) -> list[list[str]]:
    """The index as plain strings, in the order the actions were planned."""
    return [_row(action) for action in plan.actions]


def _row(action: Action) -> list[str]:
    message = action.message
    return [
        message.date.strftime("%Y-%m-%d %H:%M") if message.date else "",
        message.sender,
        message.subject,
        action.attachment.raw_filename if action.attachment else action.declared_filename,
        action.relpath,
        str(action.size) if action.attachment else "",
        action.rule,
        action.status,
        _note(action),
        message.source.name,
    ]


def _note(action: Action) -> str:
    """The action's own note, plus anything odd about the message itself.

    Header defects ride along on the row rather than getting a row of their
    own: "the Subject was undecodable" is context for a file that did get
    filed, not a separate event.
    """
    parts = [action.note, *action.message.defects]
    return "; ".join(part for part in parts if part)


def write_csv(path: Path, plan: Plan) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(COLUMNS)
        writer.writerows(rows(plan))
    return len(plan.actions)


def write_xlsx(path: Path, plan: Plan) -> int:
    """Two sheets: every row, and the counts that summarise them."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    workbook.properties.created = EPOCH
    workbook.properties.modified = EPOCH
    workbook.properties.creator = "inbox-filer"
    workbook.properties.lastModifiedBy = "inbox-filer"

    sheet = workbook.active
    sheet.title = "Index"
    sheet.append(list(COLUMNS))

    header_font = Font(bold=True)
    # One fill per status that is not the boring case. The colours are the
    # sheet saying "look at these" without the client having to read the
    # status column first.
    fills = {
        "filed-renamed": PatternFill("solid", fgColor="FFF2CC"),
        "no-attachment": PatternFill("solid", fgColor="E8E8E8"),
        "unreadable-attachment": PatternFill("solid", fgColor="FCE4E4"),
    }
    unsorted_fill = PatternFill("solid", fgColor="FCE4E4")

    size_column = COLUMNS.index("size_bytes") + 1
    for action, row in zip(plan.actions, rows(plan)):
        sheet.append(row)
        index = sheet.max_row
        cell = sheet.cell(row=index, column=size_column)
        # Written as text above so the CSV and the sheet agree; here it becomes
        # a number, because a client sorting by size expects it to sort.
        cell.value = action.size if action.attachment else None
        fill = fills.get(action.status)
        if action.rule == "unsorted":
            fill = unsorted_fill
        if fill is not None:
            for column in range(1, len(COLUMNS) + 1):
                sheet.cell(row=index, column=column).fill = fill

    for column in range(1, len(COLUMNS) + 1):
        sheet.cell(row=1, column=column).font = header_font
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{sheet.max_row}"

    widths = {"date": 17, "sender": 34, "subject": 40, "original_filename": 30,
              "new_path": 52, "size_bytes": 11, "rule": 18, "status": 21,
              "note": 60, "source_message": 22}
    for position, name in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(position)].width = widths[name]
    wrap = Alignment(vertical="top", wrap_text=True)
    for row_cells in sheet.iter_rows(min_row=2, min_col=COLUMNS.index("note") + 1,
                                     max_col=COLUMNS.index("note") + 1):
        for cell in row_cells:
            cell.alignment = wrap

    summary = workbook.create_sheet("Summary")
    summary.append(["what", "count"])
    summary.cell(row=1, column=1).font = header_font
    summary.cell(row=1, column=2).font = header_font
    for label, value in summary_rows(plan):
        summary.append([label, value])
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 10

    buffer = io.BytesIO()
    workbook.save(buffer)
    path.write_bytes(_repack(buffer.getvalue()))
    return len(plan.actions)


def summary_rows(plan: Plan) -> list[tuple[str, int]]:
    """The counts on the Summary sheet.

    Every one of them is a property of the filing cabinet. "How many did this
    run have to write" is missing on purpose: it is the number that changes
    between two runs over the same mailbox, and putting it here would make the
    sheet differ every time while nothing about the filing had changed. It is
    printed on stdout instead, where a per-run number belongs.
    """
    counts = plan.counts()
    unsorted = len(plan.unsorted)
    return [
        ("messages read", plan.message_count),
        ("attachments seen", plan.attachment_count),
        ("filed by a rule", plan.attachment_count - unsorted),
        ("of those, renamed around a collision", counts.get("filed-renamed", 0)),
        ("sent to unsorted", unsorted),
        ("attachments that would not decode", counts.get("unreadable-attachment", 0)),
        ("messages with no attachment", counts.get("no-attachment", 0)),
    ]


#: The two places an .xlsx records a wall clock. openpyxl sets `modified` to
#: the current time as it saves, whatever the workbook's properties said, so
#: pinning it has to happen here rather than before the save.
_TIMESTAMP = re.compile(
    rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:(?:created|modified)>)"
)
_EPOCH_XML = b"1980-01-01T00:00:00Z"


def _repack(data: bytes) -> bytes:
    """Rewrite an xlsx with every clock in it flattened.

    Two of them. An .xlsx is a zip, and `ZipFile.writestr` stamps each member
    with the current local time; it is also an Office document, and
    `docProps/core.xml` carries a modification timestamp that openpyxl
    refreshes on save. Either one makes two runs a second apart produce
    different bytes for identical content, which would put "run it twice and
    the index is unchanged" out of reach of `cmp` — the one check a client
    might actually perform. Order and contents are left exactly as openpyxl
    wrote them; only the clocks go.
    """
    source = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            body = source.read(item.filename)
            if item.filename == "docProps/core.xml":
                body = _TIMESTAMP.sub(rb"\g<1>" + _EPOCH_XML + rb"\g<2>", body)
            info = zipfile.ZipInfo(item.filename, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = item.compress_type
            info.external_attr = item.external_attr
            info.create_system = 0
            target.writestr(info, body)
    return out.getvalue()
